"""office 工具的 worker：按 request.json 执行一个动作，把结果写进同一调用目录的 response.json。

用法：python worker.py <call_dir>/request.json
写完 response.json 后以 0 退出，动作失败也一样；只有 worker 自身崩溃才以非 0 退出。
动作：probe、guide、read、write、view、frames、cleanup。

顶层只导入标准库与 util、com：缺第三方包时 probe 仍要能跑出缺项清单。
"""

import contextlib
import importlib.metadata
import io
import os
import runpy
import shutil
import sys
import traceback
import types
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import com  # noqa: E402
import util  # noqa: E402

# (发行包名, 导入名)。matplotlib 只在脚本用 office.figure 时需要，av 只在视频抽帧时需要，缺了都不算缺项。
PACKAGES = [("python-docx", "docx"), ("openpyxl", "openpyxl"), ("python-pptx", "pptx"),
            ("lxml", "lxml"), ("Pillow", "PIL"), ("pypdfium2", "pypdfium2"),
            ("psutil", "psutil"), ("matplotlib", "matplotlib"), ("av", "av")]
if sys.platform == "win32":
    PACKAGES.append(("pywin32", "win32com"))
OPTIONAL = {"matplotlib", "av"}
SCRIPT_OUTPUT_LIMIT = 20000
CJK_FONTS = ("Microsoft YaHei", "SimHei", "SimSun", "DengXian", "Noto Sans CJK SC",
             "Source Han Sans SC", "PingFang SC")


class Response:
    def __init__(self, action):
        self.data = {"ok": False, "action": action, "message": "", "stages": [], "files": [],
                     "text": "", "images": [], "script_output": "", "residue": [],
                     "instances": [], "errors": []}

    def stage(self, name, status, detail="", file=None):
        self.data["stages"].append({"name": name, "file": file, "status": status,
                                    "detail": detail})

    def __getitem__(self, key):
        return self.data[key]

    def __setitem__(self, key, value):
        self.data[key] = value


# ───────────────────────── probe / guide / read / cleanup ─────────────────────────

def do_probe(req, resp):
    packages, missing = {}, []
    for dist, _ in PACKAGES:
        try:
            packages[dist] = importlib.metadata.version(dist)
        except importlib.metadata.PackageNotFoundError:
            packages[dist] = None
            if dist not in OPTIONAL:
                missing.append(dist)
    apps = {f: com.default_app(f) for f in ("docx", "xlsx", "pptx")}
    resp["probe"] = {"python": ".".join(map(str, sys.version_info[:3])),
                     "executable": sys.executable, "packages": packages, "missing": missing,
                     "apps": apps, "platform": sys.platform}
    lines = [f"Python {resp['probe']['python']}（{sys.executable}）"]
    lines.append("缺少：" + "、".join(missing) if missing else "所需的包齐全")
    for f, a in apps.items():
        lines.append(f"{f}：{a['progid']} {a['version'] or ''} 规则 {a['rule']}" if a["available"]
                     else f"{f}：不可用（{a['reason']}）")
    resp["text"] = "\n".join(lines)
    resp["ok"] = not missing
    resp["message"] = "环境齐全" if not missing else f"缺少 {len(missing)} 个包"


def do_guide(req, resp):
    fmt = req.get("format")
    path = HERE / "guides" / f"{fmt}.md"
    if not path.is_file():
        resp["message"] = f"没有 {fmt} 的说明；format 取 docx / pptx / xlsx"
        return
    resp["text"] = path.read_text("utf-8")
    resp["ok"] = True
    resp["message"] = f"{fmt} 的做法要点"


def _check_readable(path, resp):
    p = Path(path)
    if not p.is_file():
        resp["message"] = f"{path} 不存在"
        return None
    fmt = util.fmt_of(p)
    if fmt is None:
        resp["message"] = f"{p.name} 不是 docx / pptx / xlsx"
        return None
    if util.is_ole(p):
        resp["message"] = (f"{p.name} 已加密，或是改了扩展名的旧版二进制文件，打不开；"
                           f"需要先在办公软件里去掉密码或另存为 .{fmt}")
        return None
    if util.has_macros(p):
        resp["message"] = f"{p.name} 含宏，不打开"
        return None
    return fmt


def _check_viewable(path, resp):
    """view 另收 PDF 原件，按页直接栅格化；read、write 的格式集不变，不要把 PDF 加进 util.FORMATS。"""
    p = Path(path)
    if not util.is_pdf(p):
        return _check_readable(path, resp)
    if not p.is_file():
        resp["message"] = f"{path} 不存在"
        return None
    return "pdf"


def do_read(req, resp):
    import structure

    path = req["path"]
    fmt = _check_readable(path, resp)
    if fmt is None:
        return
    sha = util.sha256(path)
    text, checks = structure.read(path, fmt, req.get("range"))
    resp["text"] = f"sha256 {sha}\n{text}"
    resp["files"].append({"path": path, "committed": False, "sha256": sha, "candidate": None,
                          "pages": None, "sheets": None, "checks": checks})
    resp["ok"] = True
    resp["message"] = f"已读取 {Path(path).name}"


def do_cleanup(req, resp):
    report = com.cleanup(req["call_dir"])
    resp["instances"] = report
    ended = sum(1 for r in report if r["action"] == "ended")
    left = [r for r in report if r["action"] in ("reported", "end_failed")]
    resp["ok"] = not left
    resp["message"] = (f"结束了 {ended} 个本次调用新起的办公软件进程"
                       + (f"；{len(left)} 个未结束：" + "、".join(
                           f"{r['pid']}（{r.get('status')}）" for r in left) if left else ""))


# ───────────────────────── view ─────────────────────────

def do_view(req, resp):
    import render

    path = req["path"]
    fmt = _check_viewable(path, resp)
    if fmt is None:
        return
    cur = util.sha256(path)
    expected = req.get("expected_sha256")
    if expected and cur != expected:
        resp["message"] = f"{Path(path).name} 在查看期间被改动；重新 view"
        return
    pages = [str(p) for p in (req.get("pages") or [])]
    if not pages:
        resp["message"] = "pages 不能为空"
        return
    apps = com.Apps(req["call_dir"], "view")
    try:
        manifest, images, missing = render.view(path, fmt, pages, req.get("region"),
                                                req["cache_dir"], req["call_dir"], apps)
    except (com.Unavailable, com.Busy, render.Unreadable) as e:
        resp["message"] = f"无法渲染：{e}"
        return
    finally:
        apps.close()
        resp["residue"] = apps.residue
    resp["images"] = images
    resp["files"].append({"path": path, "committed": False, "sha256": cur, "candidate": None,
                          "pages": None if fmt == "xlsx" else len(manifest["pages"]),
                          "sheets": manifest.get("sheets"), "checks": []})
    keys = "、".join(p["key"] for p in manifest["pages"][:60])
    if missing:
        resp["errors"].append(f"没有这些页：{'、'.join(missing)}；可用：{keys}")
    resp["ok"] = bool(images) and not missing
    resp["message"] = f"返回 {len(images)} 张图" + (f"，{len(missing)} 个页码不存在" if missing else "")


# ───────────────────────── write ─────────────────────────

class OfficeContext:
    """注入脚本全局变量 office 的对象；同时以 `from qyoffice import office` 可取到。"""

    def __init__(self, workspace, call_dir, inputs, works):
        self.workspace = str(workspace)
        self.call_dir = str(call_dir)
        self.inputs = [str(p) for p in inputs]
        self._works = works
        self._apps = None

    def output(self, path):
        """声明过的输出对应的工作副本路径。脚本写这个路径，worker 检查、渲染后再提交到正式路径。"""
        target = path if os.path.isabs(path) else os.path.join(self.workspace, path)
        for o, _, work in self._works:
            if util.same_path(o["path"], target):
                return str(work)
        declared = "、".join(o["path"] for o, _, _ in self._works)
        raise KeyError(f"{path} 不是声明过的输出；已声明：{declared}")

    def app(self, kind):
        """每次调用新起一个办公软件实例并登记，返回它的原生对象；脚本结束时统一回收。

        不要用它另存文件，一个实例也只打开一个文件：WPS 另存、以及 WPS 文字在同一实例里再次打开文件，
        都会写入用户的最近文档与账号打开记录。
        """
        if kind not in ("docx", "xlsx", "pptx"):
            raise ValueError("kind 取 docx / xlsx / pptx")
        if self._apps is None:
            self._apps = com.Apps(self.call_dir, "script", reuse=False)
        return self._apps.get(kind).app

    def figure(self, width_cm, height_cm, **kwargs):
        """按最终插入尺寸建 matplotlib 画布：插入文档时用同一宽度，图里的字号就是最终磅值。"""
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib import font_manager

        names = {f.name for f in font_manager.fontManager.ttflist}
        for name in CJK_FONTS:
            if name in names:
                plt.rcParams["font.sans-serif"] = [name] + list(plt.rcParams["font.sans-serif"])
                plt.rcParams["font.family"] = "sans-serif"
                break
        plt.rcParams["axes.unicode_minus"] = False
        return plt.subplots(figsize=(width_cm / 2.54, height_cm / 2.54), **kwargs)

    def close(self):
        if self._apps is None:
            return []
        self._apps.close()
        return self._apps.residue


def run_script(script, ctx):
    """执行模型的制作脚本。返回 (是否成功, 输出)。脚本所在目录放在 sys.path 最前，便于导入同目录模块。"""
    buf = io.StringIO()
    saved_path = list(sys.path)
    sys.path.insert(0, str(Path(script).parent))
    os.environ.setdefault("MPLBACKEND", "Agg")
    module = types.ModuleType("qyoffice")
    module.office = ctx
    sys.modules["qyoffice"] = module
    ok = True
    try:
        with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(buf):
            runpy.run_path(str(script), init_globals={"office": ctx}, run_name="__main__")
    except SystemExit as e:
        if e.code not in (None, 0):
            ok = False
            buf.write(f"\nSystemExit({e.code})\n")
    except BaseException:
        ok = False
        buf.write("\n" + traceback.format_exc())
    finally:
        sys.path[:] = saved_path
        sys.modules.pop("qyoffice", None)
    return ok, buf.getvalue()


def validate_write(req):
    problems = []
    script = Path(req.get("script") or "")
    if script.suffix.lower() != ".py" or not script.is_file():
        problems.append(f"脚本 {script} 不存在或不是 .py 文件")
    outputs = req.get("outputs") or []
    if not outputs:
        problems.append("outputs 不能为空")
    seen = []
    for o in outputs:
        p = o.get("path") or ""
        if util.fmt_of(p) is None:
            problems.append(f"{p} 不是 docx / pptx / xlsx（带宏的格式不支持）")
        if any(util.same_path(p, s) for s in seen):
            problems.append(f"{p} 重复声明")
        seen.append(p)
    for i in req.get("inputs") or []:
        if not Path(i).is_file():
            problems.append(f"输入 {i} 不存在")
    for p in req.get("pdf") or []:
        if not any(util.same_path(p.get("source", ""), s) for s in seen):
            problems.append(f"PDF 的来源 {p.get('source')} 不在 outputs 里")
        if Path(p.get("path", "")).suffix.lower() != ".pdf":
            problems.append(f"{p.get('path')} 不是 .pdf")
    return problems


def check_targets(outputs):
    """目标的当前字节必须与 expected_sha256 一致；expected 为 null 表示目标此时不应存在。"""
    conflicts = []
    for o in outputs:
        cur, exp = util.sha256_or_none(o["path"]), o.get("expected_sha256")
        if cur == exp:
            continue
        if exp is None:
            conflicts.append(f"{o['path']} 已存在；修改已有文件先用 read_office 读取，再写")
        elif cur is None:
            conflicts.append(f"{o['path']} 在读取之后被删除；重新 read 后再写")
        else:
            conflicts.append(f"{o['path']} 在读取之后被改动过；重新 read 后再写")
    return conflicts


def open_check(path, fmt):
    """包能被对应的库打开；打不开返回原因。"""
    try:
        if fmt == "docx":
            from docx import Document

            Document(str(path))
        elif fmt == "xlsx":
            from openpyxl import load_workbook

            load_workbook(str(path))
        else:
            from pptx import Presentation

            Presentation(str(path))
    except Exception as e:
        return f"文件打不开：{type(e).__name__}: {e}"
    return None


def postprocess(fmt, work, req, apps, call_dir, resp, file):
    import docx_toc
    import render
    import xlsx_calc

    try:
        if fmt == "xlsx":
            values = xlsx_calc.read_values(apps.get("xlsx"), work)
            rep = xlsx_calc.write_back(work, values)
            if rep["result"] == "ok":
                resp.stage("postprocess", "completed",
                           f"重算并写回 {rep['written']} 个公式结果", file)
            else:
                detail = "；".join((rep["mismatch"] + rep["unknown_errors"])[:5])
                resp.stage("postprocess", "failed", f"计算结果写回失败，文件保持无缓存值：{detail}", file)
        elif fmt == "docx":
            if not req.get("update_toc"):
                resp.stage("postprocess", "not_run", "未请求 update_toc", file)
                return
            rep = docx_toc.update(work, apps, call_dir)
            resp.stage("postprocess", rep["status"], rep["detail"], file)
            if rep.get("pdf"):
                # 最后一次核对的 PDF 由回填后的工作副本导出，与交付字节一致，直接作为渲染缓存。
                info = com.default_app("docx")
                render.seed_docx(work, rep["pdf"], req["cache_dir"],
                                 {k: info.get(k) for k in ("progid", "product", "version")})
        else:
            resp.stage("postprocess", "not_run", "pptx 不做后处理", file)
    except (com.Unavailable, com.Busy) as e:
        tail = "；交付件没有计算结果缓存，打开时由软件重算" if fmt == "xlsx" else ""
        resp.stage("postprocess", "unavailable", f"{e}{tail}", file)


def do_write(req, resp, apps_factory=com.Apps):
    import render
    import structure

    problems = validate_write(req)
    if problems:
        resp["errors"].extend(problems)
        resp["message"] = problems[0]
        resp.stage("execute", "not_run", "；".join(problems))
        return
    outputs = req["outputs"]
    conflicts = check_targets(outputs)
    if conflicts:
        resp["errors"].extend(conflicts)
        resp["message"] = conflicts[0]
        resp.stage("execute", "not_run", "；".join(conflicts))
        return

    call = Path(req["call_dir"])
    work_dir = call / "work"
    work_dir.mkdir(parents=True, exist_ok=True)
    works = []
    for i, o in enumerate(outputs):
        final = Path(o["path"])
        work = work_dir / f"{i}-{final.name}"
        if final.exists():
            shutil.copyfile(final, work)
        works.append((o, util.fmt_of(final), work))

    script = Path(req["script"])
    shutil.copyfile(script, call / "script.py")
    script_sha = util.sha256(script)
    ctx = OfficeContext(req["workspace"], call, req.get("inputs") or [], works)
    ok, output = run_script(script, ctx)
    script_residue = ctx.close()
    resp["script_output"] = util.truncate(output, SCRIPT_OUTPUT_LIMIT)
    if util.sha256(script) != script_sha:
        resp["errors"].append("脚本在执行期间被改动，调用目录里的 script.py 是执行前的版本")
    if not ok:
        resp.stage("execute", "failed", "脚本出错，输出见 script_output")
        for o, _, work in works:
            resp["files"].append({"path": o["path"], "committed": False, "sha256": None,
                                  "candidate": str(work) if work.exists() else None, "pages": None,
                                  "sheets": None, "checks": []})
        resp["message"] = "脚本出错，没有提交任何文件"
        resp["residue"] = script_residue
        resp["instances"] = com._load_instances(call)
        return
    resp.stage("execute", "completed", "脚本执行完成")

    apps = apps_factory(call, "post")
    try:
        for o, fmt, work in works:
            final = o["path"]
            entry = {"path": final, "committed": False, "sha256": None, "candidate": str(work),
                     "pages": None, "sheets": None, "checks": []}
            resp["files"].append(entry)
            if not work.exists():
                entry["candidate"] = None
                resp.stage("commit", "not_run", "脚本没有写出这个文件（写 office.output(路径) 返回的位置）",
                           final)
                continue
            err = open_check(work, fmt) or ("文件含宏" if util.has_macros(work) else None)
            if err:
                resp.stage("check", "failed", err, final)
                resp.stage("commit", "not_run", "文件未通过打开检查", final)
                continue
            postprocess(fmt, work, req, apps, call, resp, final)
            entry["checks"] = structure.check(work, fmt)
            resp.stage("check", "completed", f"{len(entry['checks'])} 条检查结果", final)
            try:
                manifest = render.render(work, fmt, req["cache_dir"], apps)
                entry["pages"] = None if fmt == "xlsx" else len(manifest["pages"])
                entry["sheets"] = manifest.get("sheets")
                resp.stage("render", "completed", f"{len(manifest['pages'])} 页已渲染", final)
            except (com.Unavailable, com.Busy) as e:
                resp.stage("render", "unavailable", str(e), final)
            except Exception as e:
                resp.stage("render", "failed", f"{type(e).__name__}: {e}", final)
            cur = util.sha256_or_none(final)
            if cur != o.get("expected_sha256"):
                resp.stage("commit", "failed", "目标在执行期间被改动，没有覆盖；重新 read 后再写", final)
                continue
            util.atomic_copy(work, final)
            entry["committed"] = True
            entry["sha256"] = util.sha256(final)
            resp.stage("commit", "completed", "", final)
        for p in req.get("pdf") or []:
            src = next((e for e in resp["files"] if util.same_path(e["path"], p["source"])), None)
            if not src or not src["committed"]:
                resp.stage("commit", "not_run", "来源文件没有提交，不导出 PDF", p["path"])
                continue
            try:
                render.export_pdf(src["path"], util.fmt_of(src["path"]), p["path"], req["cache_dir"],
                                  apps)
                resp.stage("commit", "completed", "PDF 已导出", p["path"])
            except (com.Unavailable, com.Busy) as e:
                resp.stage("commit", "unavailable", f"PDF 未导出：{e}", p["path"])
            except Exception as e:
                resp.stage("commit", "failed", f"PDF 未导出：{type(e).__name__}: {e}", p["path"])
    finally:
        apps.close()
        seen = {r["pid"]: r for r in script_residue + apps.residue}
        resp["residue"] = list(seen.values())
        resp["instances"] = com._load_instances(call)
    committed = [e for e in resp["files"] if e["committed"]]
    resp["ok"] = len(committed) == len(works)
    resp["message"] = (f"已提交 {len(committed)} / {len(works)} 个文件"
                       + ("" if resp["ok"] else "，其余见 stages"))


# ───────────────────────── frames ─────────────────────────

def do_frames(req, resp):
    """视频抽帧：带时间戳的 JPEG 与一段说明。缺 av 时如实报缺解码库，不影响其余动作。"""
    path = Path(req["path"])
    if not path.is_file():
        resp["message"] = f"{req['path']} 不存在"
        return
    import video

    try:
        import av  # noqa: F401
    except ImportError:
        resp["message"] = "缺少视频解码库 av（PyAV）"
        return
    try:
        result = video.sample(path, Path(req["call_dir"]) / "frames", req.get("start"), req.get("end"),
                              req.get("max_frames") or 12, req.get("max_edge") or 1024)
    except video.Unreadable as e:
        resp["message"] = str(e)
        return
    resp["images"] = [{"path": f["path"], "label": video.clock(f["time"]), "width": f["width"],
                       "height": f["height"]} for f in result["frames"]]
    resp["text"] = video.describe(path.name, result)
    resp["ok"] = True
    resp["message"] = f"返回 {len(resp['images'])} 帧"


HANDLERS = {"probe": do_probe, "guide": do_guide, "read": do_read, "write": do_write,
            "view": do_view, "frames": do_frames, "cleanup": do_cleanup}


def main(argv):
    req = util.read_json(argv[1])
    action = req.get("action")
    resp = Response(action)
    if sys.platform == "win32":
        try:
            import pythoncom

            pythoncom.CoInitialize()
        except ImportError:
            pass
    try:
        handler = HANDLERS.get(action)
        if handler is None:
            resp["message"] = f"未知动作 {action}"
        else:
            handler(req, resp)
    except Exception as e:
        resp["ok"] = False
        resp["errors"].append(traceback.format_exc())
        resp["message"] = resp["message"] or f"{action} 执行出错：{type(e).__name__}: {e}"
    util.write_json(Path(req["call_dir"]) / "response.json", resp.data)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
