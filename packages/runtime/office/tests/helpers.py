"""测试共用函数：把 office 目录加入导入路径，创建临时目录，生成样本文件。

临时目录取自 tempfile，受 TEMP / TMP / TMPDIR 控制；门禁把它们设置为仓库 .tmp 下的目录。
"""

import shutil
import sys
import tempfile
from pathlib import Path

OFFICE_DIR = Path(__file__).resolve().parent.parent
if str(OFFICE_DIR) not in sys.path:
    sys.path.insert(0, str(OFFICE_DIR))


class TempDir:
    def __enter__(self):
        self.path = Path(tempfile.mkdtemp(prefix="qy-office-test-"))
        return self.path

    def __exit__(self, *exc):
        shutil.rmtree(self.path, ignore_errors=True)


def toc_docx(path, prefix_title=False):
    """带一个 TOC 域（占位结果，dirty）和两级标题的 docx。"""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    p = doc.add_paragraph("目录" if prefix_title else None)

    def run_with(el):
        r = OxmlElement("w:r")
        r.append(el)
        p._p.append(r)

    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    b.set(qn("w:dirty"), "true")
    run_with(b)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = 'TOC \\o "1-2" \\h \\z \\u'
    run_with(it)
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    run_with(s)
    t = OxmlElement("w:t")
    t.text = "（目录待更新）"
    run_with(t)
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    run_with(e)
    doc.add_heading("第一章 概述", level=1)
    doc.add_paragraph("正文一。")
    doc.add_heading("1.1 背景", level=2)
    doc.add_paragraph("正文二。")
    doc.add_heading("第二章 市场", level=1)
    doc.add_heading("三级标题不进目录", level=3)
    doc.save(path)
    return path
