"""覆盖 xlsx_calc.py 的 convert 与 write_back，以及 util.py 的 replace_parts。"""

import hashlib
import unittest
import zipfile

from helpers import TempDir

import xlsx_calc
from com import XL_ERRORS


def workbook(path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 5
    ws["B1"] = "=A1*3"
    ws["B2"] = "=A1/0"
    ws["B3"] = '="x"&A1'
    ws["B4"] = "=A1>2"
    ws["B5"] = '=""'
    wb.create_sheet("空表")
    wb.save(path)
    return path


VALUES = {"Sheet1": {"B1": ("n", "15"), "B2": ("e", "#DIV/0!"), "B3": ("str", "x5"),
                     "B4": ("b", "1"), "B5": ("str", "")}, "空表": {}}


def parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}


class ConvertTest(unittest.TestCase):
    def test_types(self):
        self.assertEqual(xlsx_calc.convert(True), ("b", "1"))
        self.assertEqual(xlsx_calc.convert(False), ("b", "0"))
        self.assertEqual(xlsx_calc.convert(-2146826281), ("e", "#DIV/0!"))
        self.assertEqual(xlsx_calc.convert(-1), ("?", "-1"))
        self.assertEqual(xlsx_calc.convert(15.0), ("n", "15"))
        self.assertEqual(xlsx_calc.convert(0.25), ("n", "0.25"))
        self.assertEqual(xlsx_calc.convert("x"), ("str", "x"))
        self.assertIsNone(xlsx_calc.convert(None))
        self.assertEqual(set(XL_ERRORS.values()) >= {"#REF!", "#N/A"}, True)


class WriteBackTest(unittest.TestCase):
    def test_writes_cached_values_and_only_touches_sheet_and_workbook(self):
        from openpyxl import load_workbook

        with TempDir() as d:
            path = workbook(d / "a.xlsx")
            before = parts(path)
            rep = xlsx_calc.write_back(path, VALUES)
            self.assertEqual(rep["result"], "ok", rep)
            self.assertEqual(rep["written"], 5)
            after = parts(path)
            changed = sorted(n for n in before if before[n] != after[n])
            self.assertEqual(changed, ["xl/workbook.xml", "xl/worksheets/sheet1.xml"])
            v = load_workbook(path, data_only=True)["Sheet1"]
            self.assertEqual((v["B1"].value, v["B2"].value, v["B3"].value, v["B4"].value),
                             (15, "#DIV/0!", "x5", True))
            f = load_workbook(path)["Sheet1"]
            self.assertEqual(f["B1"].value, "=A1*3")
            with zipfile.ZipFile(path) as z:
                self.assertIn(b'fullCalcOnLoad="1"', z.read("xl/workbook.xml"))

    def test_array_formula_members_get_cached_values(self):
        from openpyxl import Workbook, load_workbook
        from openpyxl.worksheet.formula import ArrayFormula

        with TempDir() as d:
            path = d / "arr.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "S"
            for r in range(1, 4):
                ws[f"A{r}"] = r
            ws["B1"] = ArrayFormula("B1:B3", "=A1:A3*10")
            ws["D5"] = 1
            wb.save(path)
            rep = xlsx_calc.write_back(path, {"S": {"B1": ("n", "10"), "B2": ("n", "20"),
                                                    "B3": ("n", "30")}})
            self.assertEqual(rep["result"], "ok", rep)
            v = load_workbook(path, data_only=True)["S"]
            self.assertEqual([v[f"B{r}"].value for r in (1, 2, 3)], [10, 20, 30])
            self.assertEqual(v["D5"].value, 1)
            f = load_workbook(path)["S"]
            self.assertEqual(f["B1"].value.text, "=A1:A3*10")

    def test_missing_formula_value_leaves_file_unchanged(self):
        with TempDir() as d:
            path = workbook(d / "a.xlsx")
            raw = path.read_bytes()
            values = {"Sheet1": dict(VALUES["Sheet1"]), "空表": {}}
            del values["Sheet1"]["B3"]
            rep = xlsx_calc.write_back(path, values)
            self.assertEqual(rep["result"], "failed")
            self.assertTrue(any("B3" in m for m in rep["mismatch"]))
            self.assertEqual(path.read_bytes(), raw)

    def test_unknown_error_code_leaves_file_unchanged(self):
        with TempDir() as d:
            path = workbook(d / "a.xlsx")
            raw = path.read_bytes()
            values = {"Sheet1": {**VALUES["Sheet1"], "B2": ("?", "-2146820000")}, "空表": {}}
            rep = xlsx_calc.write_back(path, values)
            self.assertEqual(rep["result"], "failed")
            self.assertTrue(rep["unknown_errors"])
            self.assertEqual(path.read_bytes(), raw)


if __name__ == "__main__":
    unittest.main()
