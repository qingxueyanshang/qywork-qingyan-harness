"""覆盖 structure.py 的三种格式检查（check_docx / check_pptx / check_xlsx / check_metadata）、同类合并与结构提取（read_*）。"""

import unittest

from helpers import TempDir, toc_docx

import structure
import xlsx_calc


def codes(checks):
    return [c["code"] for c in checks]


class DocxTest(unittest.TestCase):
    def test_east_asia_font(self):
        from docx import Document
        from docx.oxml.ns import qn

        with TempDir() as d:
            doc = Document()
            doc.add_paragraph("中文")
            doc.save(d / "a.docx")
            self.assertIn("east_asia_font_missing", codes(structure.check_docx(d / "a.docx")))
            doc.styles["Normal"].element.get_or_add_rPr().get_or_add_rFonts().set(
                qn("w:eastAsia"), "宋体")
            doc.save(d / "b.docx")
            self.assertNotIn("east_asia_font_missing", codes(structure.check_docx(d / "b.docx")))

    def test_image_wider_than_text(self):
        import io

        from docx import Document
        from docx.shared import Cm
        from PIL import Image

        buf = io.BytesIO()
        Image.new("RGB", (40, 20), "red").save(buf, "PNG")
        with TempDir() as d:
            doc = Document()
            buf.seek(0)
            doc.add_picture(buf, width=Cm(25))
            doc.save(d / "a.docx")
            self.assertIn("image_wider_than_text", codes(structure.check_docx(d / "a.docx")))

    def test_toc_not_updated_and_read(self):
        with TempDir() as d:
            path = toc_docx(d / "a.docx")
            self.assertIn("toc_not_updated", codes(structure.check_docx(path)))
            text = structure.read_docx(path)
            self.assertIn("第一章 概述", text)
            self.assertIn("域 TOC", text)


    def test_update_fields_on_open(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        with TempDir() as d:
            doc = Document()
            doc.save(d / "a.docx")
            self.assertNotIn("update_fields_on_open", codes(structure.check_docx(d / "a.docx")))
            el = OxmlElement("w:updateFields")
            el.set(qn("w:val"), "true")
            doc.settings.element.append(el)
            doc.save(d / "b.docx")
            self.assertIn("update_fields_on_open", codes(structure.check_docx(d / "b.docx")))


    def test_range_accepts_paragraph_labels(self):
        self.assertEqual(structure._parse_range("P62-P72"), (62, 72))
        self.assertEqual(structure._parse_range("3-5"), (3, 5))
        self.assertEqual(structure._parse_range("p7"), (7, 7))


class PptxTest(unittest.TestCase):
    def test_off_slide_overlap_placeholder_notes(self):
        from pptx import Presentation
        from pptx.util import Inches

        with TempDir() as d:
            prs = Presentation()
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = "标题"
            a = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
            a.text_frame.text = "互相遮挡的文本框甲"
            b = s.shapes.add_textbox(Inches(2), Inches(1.1), Inches(4), Inches(2))
            b.text_frame.text = "互相遮挡的文本框乙"
            # 外框大面积重叠，文字位于各自外框的不同角落：不算遮挡。
            e = s.shapes.add_textbox(Inches(1), Inches(4), Inches(5), Inches(2))
            e.text_frame.text = "左上"
            f = s.shapes.add_textbox(Inches(2), Inches(4.5), Inches(5), Inches(2))
            f.text_frame.text = "另一段"
            f.text_frame.paragraphs[0].alignment = 3  # 右对齐
            c = s.shapes.add_textbox(Inches(9), Inches(1), Inches(3), Inches(1))
            c.text_frame.text = "越界"
            s.notes_slide.notes_text_frame.text = "备注"
            prs.save(d / "a.pptx")
            got = codes(structure.check_pptx(d / "a.pptx"))
            for code in ("off_slide", "text_overlap", "empty_placeholder", "notes_present"):
                self.assertIn(code, got)
            overlaps = [c["message"] for c in structure.check_pptx(d / "a.pptx") if c["code"] == "text_overlap"]
            self.assertEqual(len(overlaps), 1)
            self.assertIn(a.name, overlaps[0])
            text = structure.read_pptx(d / "a.pptx")
            self.assertIn("越界", text)
            self.assertIn("备注", text)

    def test_read_lists_group_members_at_slide_position(self):
        from PIL import Image
        from pptx import Presentation
        from pptx.oxml.ns import qn
        from pptx.util import Inches

        with TempDir() as d:
            Image.new("RGB", (40, 20), (200, 30, 30)).save(d / "pill.png")
            prs = Presentation()
            s = prs.slides.add_slide(prs.slide_layouts[6])
            outer = s.shapes.add_group_shape()
            outer.shapes.add_picture(str(d / "pill.png"), Inches(1), Inches(1), Inches(2), Inches(1))
            inner = outer.shapes.add_group_shape()
            label = inner.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(1.6), Inches(0.6))
            label.text_frame.text = "种质库"
            # 外层组整体右移 2.54 cm（子坐标框不变）：组内形状在页面上随之右移。
            off = outer._element.grpSpPr.find(qn("a:xfrm")).find(qn("a:off"))
            off.set("x", str(int(off.get("x")) + Inches(1)))
            prs.save(d / "a.pptx")
            text = structure.read_pptx(d / "a.pptx")
            self.assertIn("种质库", text)
            self.assertIn("图片 ppt/media/", text)
            self.assertIn("40×20 px", text)
            line = next(x for x in text.splitlines() if "种质库" in x)
            self.assertTrue(line.startswith("      - "), line)          # 第二层组内：缩进三级
            self.assertIn("@(5.59,3.05)", line)                         # (1.2 in + 1 in, 1.2 in)
            # 组内文字同样检查：中文没有东亚字体
            ea = [c["message"] for c in structure.check_pptx(d / "a.pptx") if c["code"] == "east_asia_font_missing"]
            self.assertTrue(ea and label.name in ea[0], ea)

    def test_number_and_unit_side_by_side_are_not_overlap(self):
        """数值与单位各占一个文本框、同一基线左右相接（精创 P2「36.92」「亿元」，坐标取自原文件）：估算的文字区域
        伸进单位的框里不到半个字，不报告。上下两行错开（精创 P12 眉题压标题）、同一基线一行压在另一行上，照常报告。"""
        from pptx import Presentation
        from pptx.util import Emu, Pt

        def box(s, name, text, x, y, w, h, size):
            t = s.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
            t.name = name
            f = t.text_frame
            f.word_wrap = False
            f.margin_left = f.margin_right = f.margin_top = f.margin_bottom = 0
            f.text = text
            f.paragraphs[0].runs[0].font.size = Pt(size)

        with TempDir() as d:
            prs = Presentation()
            s = prs.slides.add_slide(prs.slide_layouts[6])
            box(s, "数值", "36.92", 4280802, 1363691, 1757296, 905627, 58.44)
            box(s, "单位", "亿元", 6021837, 1759483, 568751, 331042, 20.76)
            box(s, "眉题", "风险防控", 684252, 230271, 831252, 251563, 15.53)
            box(s, "标题", "五类风险与应对措施", 321928, 400022, 3930217, 522085, 33.28)
            box(s, "原文", "年产值 3.2 亿元", 6846308, 4323312, 2000000, 234063, 14.39)
            box(s, "副本", "年产值 3.2 亿元", 7146308, 4323312, 2000000, 234063, 14.39)
            prs.save(d / "a.pptx")
            overlaps = [c["message"] for c in structure.check_pptx(d / "a.pptx") if c["code"] == "text_overlap"]
            self.assertEqual(len(overlaps), 2, overlaps)
            self.assertTrue(any("眉题" in m and "标题" in m for m in overlaps), overlaps)
            self.assertTrue(any("原文" in m and "副本" in m for m in overlaps), overlaps)

    def test_text_overflow_and_east_asia_font(self):
        from pptx import Presentation
        from pptx.oxml.ns import qn
        from pptx.util import Inches, Pt

        with TempDir() as d:
            prs = Presentation()
            s = prs.slides.add_slide(prs.slide_layouts[6])
            full = s.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(3), Inches(0.8))
            full.text_frame.word_wrap = True
            full.text_frame.text = "这一段文字远多于文本框能容纳的行数，" * 8
            full.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
            roomy = s.shapes.add_textbox(Inches(5), Inches(0.5), Inches(4), Inches(2))
            roomy.text_frame.word_wrap = True
            roomy.text_frame.text = "容量充足。"
            prs.save(d / "a.pptx")
            checks = structure.check_pptx(d / "a.pptx")
            over = [c["message"] for c in checks if c["code"] == "text_overflow"]
            self.assertEqual(len(over), 1)
            self.assertIn(full.name, over[0])
            self.assertIn("east_asia_font_missing", codes(checks))

            for shp in (full, roomy):
                rpr = shp.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
                rpr.append(rpr.makeelement(qn("a:ea"), {"typeface": "微软雅黑"}))
            prs.save(d / "b.pptx")
            self.assertNotIn("east_asia_font_missing", codes(structure.check_pptx(d / "b.pptx")))


class CondenseTest(unittest.TestCase):
    def test_many_of_one_code_are_summarized(self):
        checks = [{"level": "info", "code": "notes_present", "message": f"第 {i} 页有演讲者备注"}
                  for i in range(1, 21)]
        out = structure._condense(checks + [{"level": "warning", "code": "off_slide", "message": "x"}])
        notes = [c for c in out if c["code"] == "notes_present"]
        self.assertEqual(len(notes), structure.SHOWN_PER_CODE + 1)
        self.assertIn("另有 15 处", notes[-1]["message"])
        self.assertIn("off_slide", codes(out))


class MetadataTest(unittest.TestCase):
    def test_library_defaults_are_reported(self):
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation

        with TempDir() as d:
            Document().save(d / "a.docx")
            Presentation().save(d / "a.pptx")
            Workbook().save(d / "a.xlsx")
            for name in ("a.docx", "a.pptx", "a.xlsx"):
                self.assertIn("library_metadata", codes(structure.check_metadata(d / name)), name)
            doc = Document()
            doc.core_properties.author = "林晓"
            doc.core_properties.comments = ""
            doc.save(d / "b.docx")
            self.assertEqual(structure.check_metadata(d / "b.docx"), [])


class XlsxTest(unittest.TestCase):
    def test_errors_missing_cache_volatile(self):
        from openpyxl import Workbook

        with TempDir() as d:
            wb = Workbook()
            ws = wb.active
            ws.title = "S"
            ws["A1"] = 1
            ws["B1"] = "=A1/0"
            ws["B2"] = "=TODAY()"
            path = d / "a.xlsx"
            wb.save(path)
            got = codes(structure.check_xlsx(path))
            self.assertIn("missing_cache", got)
            self.assertIn("volatile_functions", got)
            xlsx_calc.write_back(path, {"S": {"B1": ("e", "#DIV/0!"), "B2": ("n", "46000")}})
            checks = structure.check_xlsx(path)
            self.assertIn("cached_errors", codes(checks))
            self.assertIn("S!B1", [c for c in checks if c["code"] == "cached_errors"][0]["message"])
            text = structure.read_xlsx(path, "S")
            self.assertIn("B1==A1/0 → '#DIV/0!'", text)


if __name__ == "__main__":
    unittest.main()
