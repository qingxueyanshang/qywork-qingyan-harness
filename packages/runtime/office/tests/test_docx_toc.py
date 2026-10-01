"""覆盖 docx_toc.py 的 find_toc、toc_levels、style_levels、headings_after、locate、rewrite_toc。"""

import unittest
import zipfile

from helpers import TempDir, toc_docx
from lxml import etree

import docx_toc

W = docx_toc.W


def load(path):
    with zipfile.ZipFile(path) as z:
        return z.read("word/document.xml"), etree.fromstring(z.read("word/styles.xml"))


def texts(p):
    return "".join(t.text or "" for t in p.iter(f"{{{W}}}t"))


class FindTest(unittest.TestCase):
    def test_find_levels_and_headings(self):
        with TempDir() as d:
            doc_xml, styles = load(toc_docx(d / "a.docx"))
        body = etree.fromstring(doc_xml).find(f"{{{W}}}body")
        toc = docx_toc.find_toc(body)
        self.assertIsNotNone(toc)
        self.assertIn("TOC", toc["instr"])
        self.assertEqual(docx_toc.toc_levels(toc["instr"]), (1, 2))
        self.assertEqual(docx_toc.toc_levels("TOC \\h"), (1, 3))
        levels = docx_toc.style_levels(styles)
        self.assertEqual(levels.get("Heading1"), 1)
        heads = docx_toc.headings_after(body, toc["end"], (1, 2), levels)
        self.assertEqual(heads, [(1, "第一章 概述"), (2, "1.1 背景"), (1, "第二章 市场")])


class LocateTest(unittest.TestCase):
    HEADS = [(1, "第一章 概述"), (2, "1.1 背景"), (1, "第二章 市场")]

    def pages(self, heads, pages):
        return [f and f[0] for f in docx_toc.locate(heads, pages)]

    def test_skips_toc_entries_and_keeps_order(self):
        pages = ["目录\r\n第一章 概述 .......... 2\r\n1.1 背景 2\r\n第二章 市场 3",
                 "第一章 概述\r\n正文一。\r\n1.1 背景\r\n正文二。",
                 "第二章 市场\r\n正文"]
        self.assertEqual(self.pages(self.HEADS, pages), [2, 2, 3])

    def test_heading_text_inside_a_sentence_is_not_the_heading(self):
        pages = ["第一章 概述 正文", "1.1 背景 正文"]
        self.assertEqual(self.pages(self.HEADS, pages), [None, None, None])

    def test_next_line_starting_with_digits(self):
        heads = [(2, "1.2 细分指标"), (3, "1.2.1 指标口径")]
        pages = ["目录\r\n（目录）", "1.2 细分指标\r\n1.2.1 指标口径\r\n正文"]
        self.assertEqual(self.pages(heads, pages), [2, 2])

    def test_duplicate_headings_follow_document_order(self):
        heads = [(1, "第1章 甲"), (2, "概述"), (1, "第2章 乙"), (2, "概述")]
        pages = ["第 1 章 甲\r\n概述\r\n第 1 章概述正文。",
                 "正文。\r\n第 2 章 乙\r\n第 2 章概述正文里有概述二字。",
                 "概述\r\n第 2 章概述正文。"]
        self.assertEqual(self.pages(heads, pages), [1, 1, 2, 3])

    def test_heading_joined_to_previous_line(self):
        heads = [(1, "第4章 业务单元4"), (2, "概述")]
        pages = ["数据来源说明。数据来源说明。 第 4 章 业务单元 4\r\n概述\r\n正文"]
        self.assertEqual(self.pages(heads, pages), [1, 1])

    def test_automatic_numbering_is_kept_in_the_entry(self):
        heads = [(1, "概述"), (2, "背景")]
        pages = ["一、概述\r\n正文。\r\n1.1 背景\r\n正文。"]
        self.assertEqual(docx_toc.locate(heads, pages), [(1, "一、 概述"), (1, "1.1 背景")])

    def test_hidden_heading_is_not_listed(self):
        from docx import Document

        with TempDir() as d:
            path = toc_docx(d / "a.docx")
            doc = Document(path)
            h = doc.add_heading("隐藏的标题", level=1)
            for r in h.runs:
                r.font.hidden = True
            doc.save(path)
            doc_xml, styles = load(path)
        body = etree.fromstring(doc_xml).find(f"{{{W}}}body")
        toc = docx_toc.find_toc(body)
        heads = docx_toc.headings_after(body, toc["end"], (1, 2), docx_toc.style_levels(styles))
        self.assertNotIn("隐藏的标题", [t for _, t in heads])


class RewriteTest(unittest.TestCase):
    ENTRIES = [(1, "第一章 概述", 2), (2, "1.1 背景", 2), (1, "第二章 市场", 3)]

    def check(self, new_xml, expect_head=None):
        body = etree.fromstring(new_xml).find(f"{{{W}}}body")
        toc = docx_toc.find_toc(body)
        self.assertIsNotNone(toc, "域结构被破坏")
        self.assertIn('TOC \\o "1-2"', toc["instr"])
        self.assertIsNone(toc["begin"].find(f"{{{W}}}fldChar").get(f"{{{W}}}dirty"))
        paras = list(body)
        first = paras.index(toc["begin"].getparent())
        last = paras.index(toc["end"].getparent())
        entries = [texts(p) for p in paras[first:last + 1]]
        self.assertEqual(entries, ["第一章 概述2", "1.1 背景2", "第二章 市场3"])
        tab = paras[first].find(f".//{{{W}}}tabs/{{{W}}}tab")
        self.assertEqual((tab.get(f"{{{W}}}val"), tab.get(f"{{{W}}}leader")), ("right", "dot"))
        if expect_head:
            self.assertEqual(texts(paras[first - 1]), expect_head)
        heading_after = [texts(p) for p in paras[last + 1:] if texts(p)]
        self.assertEqual(heading_after[0], "第一章 概述")

    def test_single_paragraph_field(self):
        with TempDir() as d:
            doc_xml, _ = load(toc_docx(d / "a.docx"))
        self.check(docx_toc.rewrite_toc(doc_xml, self.ENTRIES, set()))

    def test_title_before_field_in_same_paragraph_is_kept(self):
        with TempDir() as d:
            doc_xml, _ = load(toc_docx(d / "a.docx", prefix_title=True))
        self.check(docx_toc.rewrite_toc(doc_xml, self.ENTRIES, set()), expect_head="目录")

    def test_multi_paragraph_field_is_rewritten_again(self):
        with TempDir() as d:
            doc_xml, _ = load(toc_docx(d / "a.docx"))
        once = docx_toc.rewrite_toc(doc_xml, self.ENTRIES, {1, 2})
        again = docx_toc.rewrite_toc(once, [(1, "第一章 概述", 3), (2, "1.1 背景", 3),
                                            (1, "第二章 市场", 4)], {1, 2})
        body = etree.fromstring(again).find(f"{{{W}}}body")
        toc = docx_toc.find_toc(body)
        paras = list(body)
        span = paras[paras.index(toc["begin"].getparent()): paras.index(toc["end"].getparent()) + 1]
        self.assertEqual([texts(p) for p in span], ["第一章 概述3", "1.1 背景3", "第二章 市场4"])
        self.assertEqual(span[0].find(f".//{{{W}}}pStyle").get(f"{{{W}}}val"), "TOC1")

    def test_no_toc_raises(self):
        from docx import Document

        with TempDir() as d:
            Document().save(d / "b.docx")
            doc_xml, _ = load(d / "b.docx")
        with self.assertRaises(ValueError):
            docx_toc.rewrite_toc(doc_xml, self.ENTRIES, set())


if __name__ == "__main__":
    unittest.main()
